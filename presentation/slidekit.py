"""
slidekit -- a small component library for clean data-analysis slides.

Everything here is a *component*: a function that draws one reusable block onto a
slide. Build a deck by calling them. Nothing in this file is specific to the
Effective Engine talk -- see build_deck.py for that.

The visual identity mirrors presentation/talk.tex (midnight blue + pine green),
so the PowerPoint and the Beamer deck read as the same project.

    python3 -c "import slidekit"      # no side effects; safe to import

Quick start:

    from slidekit import *
    d = Deck()
    s = d.slide()
    title(s, "My finding", "one-line standfirst")
    kpi_row(s, [("1.30", "Spread Sharpe"), ("0.084", "Rank IC")])
    d.save("out.pptx")
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------------------------------------------------------------------------
# Palette -- carried over from talk.tex so both decks match.
# ---------------------------------------------------------------------------
NAVY  = RGBColor(0x19, 0x19, 0x70)   # midnightblue -- primary, dominant
INK   = RGBColor(0x0F, 0x16, 0x33)   # dark slide background
PINE  = RGBColor(0x01, 0x79, 0x6F)   # pinegreen -- secondary
CLAY  = RGBColor(0xB8, 0x50, 0x42)   # sharp accent: the thing to look at
SLATE = RGBColor(0x5A, 0x64, 0x72)   # muted body / captions
RULE  = RGBColor(0xD3, 0xDA, 0xE6)   # hairlines, holder outlines
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

TINT_NAVY = RGBColor(0xEE, 0xF1, 0xF8)   # "block"      in talk.tex
TINT_PINE = RGBColor(0xE4, 0xF0, 0xEE)   # "alertblock" in talk.tex
TINT_CLAY = RGBColor(0xF9, 0xEE, 0xEC)   # emphasis
TINT_GREY = RGBColor(0xF4, 0xF6, 0xF9)   # neutral / placeholder

# Semantic aliases, so slide code reads in the language of the content.
PURGE_C   = RGBColor(0x9B, 0xA6, 0xB8)   # dropped: label overlaps test
EMBARGO_C = RGBColor(0xD3, 0xDA, 0xE6)   # dropped: serial correlation

TRAIN, VALID, TEST = NAVY, PINE, CLAY
PURGE, EMBARGO = PURGE_C, EMBARGO_C

# Fonts: both ship with Office *and* render true-to-width under LibreOffice,
# so the QA render can be trusted for overflow.
HEAD = "Cambria"
BODY = "Calibri"
MONO = "Courier New"

# Canvas
W, H = 13.333, 7.5
M = 0.62                      # page margin
CW = W - 2 * M                # content width
BOTTOM = 6.56                 # baseline: nothing sits below this


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self._blank = self.prs.slide_layouts[6]

    def slide(self, bg=None):
        s = self.prs.slides.add_slide(self._blank)
        if bg is not None:
            fill = s.background.fill
            fill.solid()
            fill.fore_color.rgb = bg
        return s

    def save(self, path):
        self.prs.save(path)
        return path


def _noline(shape):
    shape.line.fill.background()
    return shape


def rect(s, x, y, w, h, fill=None, line=None, lw=1.0, radius=None):
    """A rectangle. Pass radius (0-0.5) for rounded corners."""
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        _noline(sh)
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if sh.has_text_frame:
        sh.text_frame.text = ""
    return sh


def text(s, x, y, w, h, runs, size=13, color=BLACK, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False,
         space_after=4, line_spacing=None, wrap=True, shrink=False):
    """
    Text box. `runs` is a string, a list of lines, or a list of paragraph specs:
        "one line"
        ["line one", "line two"]
        [{"t": "Label: ", "b": True}, {"t": "value"}]        <- one paragraph
        [[{"t": "a", "b": True}, {"t": "b"}], "second line"] <- two paragraphs
    """
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if shrink:
        from pptx.enum.text import MSO_AUTO_SIZE
        tf.auto_size = MSO_AUTO_SIZE.NONE

    paras = runs if isinstance(runs, list) else [runs]
    if paras and isinstance(paras[0], dict):
        paras = [paras]                      # a single mixed-run paragraph

    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        specs = para if isinstance(para, list) else [{"t": para}]
        if isinstance(specs, dict):
            specs = [specs]
        for spec in specs:
            if isinstance(spec, str):
                spec = {"t": spec}
            r = p.add_run()
            r.text = spec.get("t", "")
            f = r.font
            f.size = Pt(spec.get("size", size))
            f.bold = spec.get("b", bold)
            f.italic = spec.get("i", italic)
            f.name = spec.get("font", font)
            f.color.rgb = spec.get("c", color)
    return box


def bullets(s, x, y, w, h, items, size=13, color=BLACK, gap=7, dot=NAVY,
            dot_size=0.062, indent=0.20, line_spacing=0.98):
    """
    Bulleted list drawn with real dots, so spacing is ours rather than the
    master's. `items` may be strings or [{"t":..,"b":True}, ..] run lists.
    Returns the y coordinate just past the last item.
    """
    cy = y
    for item in items:
        body_w = w - indent
        est = _wrapped_lines(item, body_w, size)
        line_h = size * line_spacing / 72.0
        blk = est * line_h
        d = dot_size
        rect(s, x + 0.035, cy + line_h / 2 - d / 2, d, d, fill=dot, radius=0.5)
        text(s, x + indent, cy, body_w, blk + 0.06, [item], size=size,
             color=color, space_after=0, line_spacing=line_spacing)
        cy += blk + gap / 72.0
    return cy


def _plain(item):
    if isinstance(item, str):
        return item
    if isinstance(item, dict):
        return item.get("t", "")
    return "".join(_plain(x) for x in item)


def _wrapped_lines(item, w_in, size):
    """Rough line count for layout. Calibri averages ~0.48em per glyph."""
    chars_per_line = max(1, int(w_in * 72 / (size * 0.48)))
    words, lines, cur = _plain(item).split(), 1, 0
    for word in words:
        add = len(word) + (1 if cur else 0)
        if cur + add > chars_per_line and cur:
            lines += 1
            cur = len(word)
        else:
            cur += add
    return lines


def line(s, x1, y1, x2, y2, color=SLATE, lw=1.25, dash=False):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(lw)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return ln


def arrow(s, x1, y1, x2, y2, color=SLATE, lw=1.5):
    """Straight arrow with a real arrowhead."""
    ln = line(s, x1, y1, x2, y2, color=color, lw=lw)
    tail = ln.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    head = tail.makeelement(qn("a:tailEnd"), {"type": "triangle",
                                              "w": "med", "len": "med"})
    tail.append(head)
    return ln


# ---------------------------------------------------------------------------
# Page furniture
# ---------------------------------------------------------------------------
def title(s, head, standfirst=None, color=NAVY):
    """Slide title, with an optional one-line standfirst underneath."""
    text(s, M, 0.40, CW, 0.62, head, size=29, bold=True, color=color, font=HEAD)
    if standfirst:
        text(s, M, 1.03, CW, 0.34, standfirst, size=13.5, color=SLATE,
             italic=True)
    return 1.50 if standfirst else 1.22


def footer(s, left, right=None, color=SLATE):
    text(s, M, H - 0.44, CW * 0.7, 0.28, left, size=9.5, color=color)
    if right:
        text(s, M + CW * 0.7, H - 0.44, CW * 0.3, 0.28, right, size=9.5,
             color=color, align=PP_ALIGN.RIGHT)


def caption(s, y, txt, w=None, x=None):
    text(s, x if x is not None else M, y, w or CW, 0.42, txt, size=10.5,
         color=SLATE, align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------------------
# Cards -- the workhorse container
# ---------------------------------------------------------------------------
def card(s, x, y, w, h, head=None, body=None, tint=TINT_NAVY, accent=NAVY,
         size=12.5, head_size=13, pad=0.22, bullet_body=True):
    """
    Tinted rounded panel with an optional bold header. `body` is a list of
    bullet items, or a string for a plain paragraph.
    """
    rect(s, x, y, w, h, fill=tint, radius=0.055)
    cy = y + pad
    if head:
        text(s, x + pad, cy, w - 2 * pad, 0.30, head, size=head_size,
             bold=True, color=accent, font=HEAD)
        cy += 0.36
    if body:
        if isinstance(body, str):
            text(s, x + pad, cy, w - 2 * pad, h - (cy - y) - pad, body,
                 size=size, color=BLACK, line_spacing=1.02)
        elif bullet_body:
            bullets(s, x + pad, cy, w - 2 * pad, h - (cy - y) - pad, body,
                    size=size, dot=accent)
        else:
            text(s, x + pad, cy, w - 2 * pad, h - (cy - y) - pad, body,
                 size=size, color=BLACK, line_spacing=1.02)
    return y + h


def note(s, x, y, w, h, head, body, accent=PINE):
    """A card in the accent colour -- talk.tex's \\begin{alertblock}."""
    tint = {NAVY: TINT_NAVY, PINE: TINT_PINE, CLAY: TINT_CLAY}.get(accent,
                                                                   TINT_PINE)
    return card(s, x, y, w, h, head, body, tint=tint, accent=accent)


def note_bottom(s, head, body, h=0.94, accent=PINE, x=M, w=None):
    """A closing note pinned to the page baseline."""
    return note(s, x, BOTTOM - h, w or CW, h, head, body, accent=accent)


def formula(s, x, y, w, txt, size=17, h=0.62, tint=TINT_NAVY, color=NAVY):
    """A display equation, centred in its own tinted band."""
    rect(s, x, y, w, h, fill=tint, radius=0.05)
    text(s, x + 0.1, y, w - 0.2, h, txt, size=size, color=color, font=HEAD,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, italic=True)
    return y + h


# ---------------------------------------------------------------------------
# COMPONENT -- stat tiles
# ---------------------------------------------------------------------------
def kpi_row(s, items, y=1.75, x=M, w=None, h=1.32, gap=0.22, accent=NAVY,
            tint=TINT_NAVY, value_size=30):
    """
    A row of big-number tiles. items = [(value, label), ...] or
    [(value, label, accent_colour), ...].
    """
    w = w or CW
    n = len(items)
    tw = (w - gap * (n - 1)) / n
    for i, item in enumerate(items):
        val, lab = item[0], item[1]
        acc = item[2] if len(item) > 2 else accent
        tnt = {NAVY: TINT_NAVY, PINE: TINT_PINE, CLAY: TINT_CLAY}.get(acc, tint)
        tx = x + i * (tw + gap)
        rect(s, tx, y, tw, h, fill=tnt, radius=0.07)
        text(s, tx + 0.12, y + 0.20, tw - 0.24, 0.56, val, size=value_size,
             bold=True, color=acc, font=HEAD, align=PP_ALIGN.CENTER)
        text(s, tx + 0.12, y + h - 0.48, tw - 0.24, 0.40, lab, size=10.5,
             color=SLATE, align=PP_ALIGN.CENTER, line_spacing=0.95)
    return y + h


# ---------------------------------------------------------------------------
# COMPONENT -- time-split geometry
# ---------------------------------------------------------------------------
def timebar(s, x, y, w, h, segments, label=None, label_w=1.35,
            label_size=10.5, seg_size=9.5, radius=None, outline=False):
    """
    One horizontal time axis split into coloured segments.

    segments = [(fraction, colour, caption), ...]; fractions need not sum to 1,
    they are normalised. caption "" draws the block with no text. This is the
    atom the split diagrams below are built from.
    """
    if label:
        text(s, x, y + h / 2 - 0.11, label_w - 0.10, 0.24, label,
             size=label_size, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
        x += label_w
        w -= label_w
    total = sum(seg[0] for seg in segments) or 1.0
    cx = x
    for frac, col, cap in segments:
        sw = w * frac / total
        if col is not None:
            sh = rect(s, cx, y, sw, h, fill=col, radius=radius,
                      line=RULE if outline else None, lw=0.75)
            if cap:
                light = col in (NAVY, PINE, CLAY, INK)
                text(s, cx, y, sw, h, cap, size=seg_size, bold=True,
                     color=WHITE if light else SLATE, align=PP_ALIGN.CENTER,
                     anchor=MSO_ANCHOR.MIDDLE)
        elif cap:
            text(s, cx, y, sw, h, cap, size=seg_size, color=SLATE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, italic=True)
        cx += sw
    return y + h


def time_axis(s, x, y, w, left="past", right="now", color=SLATE, ticks=0):
    """A thin arrow standing for the direction of time, with end labels."""
    arrow(s, x, y, x + w, y, color=color, lw=1.1)
    if ticks:
        for i in range(1, ticks):
            tx = x + w * i / ticks
            line(s, tx, y - 0.05, tx, y + 0.05, color=RULE, lw=0.9)
    text(s, x, y + 0.08, w * 0.5, 0.26, left, size=9.5, color=color)
    text(s, x + w * 0.5, y + 0.08, w * 0.5, 0.26, right, size=9.5, color=color,
         align=PP_ALIGN.RIGHT)
    return y + 0.36


def legend(s, x, y, items, size=10, box=0.145, gap=0.30, dy=0.0):
    """items = [(colour, label), ...]. Laid out on one row."""
    cx = x
    for col, lab in items:
        rect(s, cx, y + dy, box, box, fill=col, radius=0.25)
        wlab = len(lab) * size * 0.0075 + 0.10
        text(s, cx + box + 0.09, y + dy - 0.045, wlab, 0.26, lab, size=size,
             color=SLATE, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        cx += box + 0.09 + wlab + gap
    return y + 0.34


def split_diagram(s, x, y, w, rows, row_h=0.30, row_gap=0.11, label_w=1.35,
                  outline=False):
    """
    Stack of timebars -- one per fold. rows = [(label, segments), ...].
    This is the geometry behind every cross-validation scheme below.
    """
    cy = y
    for lab, segs in rows:
        timebar(s, x, cy, w, row_h, segs, label=lab, label_w=label_w,
                outline=outline)
        cy += row_h + row_gap
    return cy - row_gap


# --- named schemes: call these directly, or copy and edit the fractions -----

def split_holdout(s, x, y, w, train=0.6, val=0.2, test=0.2, row_h=0.46):
    return timebar(s, x, y, w, row_h,
                   [(train, TRAIN, "Train"), (val, VALID, "Validation"),
                    (test, TEST, "Test")], radius=0.06)


def split_walk_forward(s, x, y, w, folds=5, **kw):
    """Expanding window -- the scheme the walk-forward retrainer uses."""
    rows = []
    for i in range(folds):
        tr = 0.30 + i * 0.13
        rows.append((f"Window {i + 1}",
                     [(tr, TRAIN, "train" if i == 0 else ""),
                      (0.11, TEST, "deploy" if i == 0 else ""),
                      (max(0.001, 1 - tr - 0.11), None, "")]))
    return split_diagram(s, x, y, w, rows, **kw)


def split_rolling(s, x, y, w, folds=5, width=0.34, **kw):
    """Sliding window of fixed length -- drops the distant past."""
    rows = []
    step = (1 - width - 0.11) / max(1, folds - 1)
    for i in range(folds):
        lead = i * step
        rows.append((f"Window {i + 1}",
                     [(max(0.001, lead), None, ""),
                      (width, TRAIN, "train" if i == 0 else ""),
                      (0.11, TEST, "deploy" if i == 0 else ""),
                      (max(0.001, 1 - lead - width - 0.11), None, "")]))
    return split_diagram(s, x, y, w, rows, **kw)


def split_purged_kfold(s, x, y, w, folds=5, purge=0.045, embargo=0.045, **kw):
    """
    Purged K-fold with an embargo: the bars either side of the test block are
    dropped from training, so a label that overlaps the test window cannot leak.
    """
    rows = []
    blk = (1 - (folds - 1) * 0.0) / folds
    for i in range(folds):
        segs, pos = [], 0.0
        start = i * blk
        if start > 0:
            segs.append((start - purge if i else start, TRAIN,
                         "train" if i else ""))
            if i:
                segs.append((purge, PURGE, ""))
        segs.append((blk, TEST, "test" if i == 0 else ""))
        tail = 1 - start - blk
        if tail > 0:
            segs.append((embargo, EMBARGO, ""))
            segs.append((max(0.001, tail - embargo), TRAIN,
                         "train" if i else ""))
        rows.append((f"Fold {i + 1}", segs))
    return split_diagram(s, x, y, w, rows, outline=True, **kw)


def split_grouped(s, x, y, w, groups=6, held=(1, 4), **kw):
    """
    Group / date-block splitting: whole dates move together, so the same day
    never sits on both sides of the split. The cross-sectional scheme.
    """
    rows = []
    for r, name in enumerate(["Split 1", "Split 2"]):
        segs = []
        for g in range(groups):
            is_test = (g % 3 == r % 3)
            segs.append((1.0, TEST if is_test else TRAIN,
                         "test" if is_test else "train"))
        rows.append((name, segs))
    return split_diagram(s, x, y, w, rows, outline=True, **kw)


# ---------------------------------------------------------------------------
# COMPONENT -- plot holders
# ---------------------------------------------------------------------------
def plot_holder(s, x, y, w, h, label="chart", hint=None, image=None,
                tint=TINT_GREY):
    """
    A slot for a figure. Drop `image=path` in and it renders the real chart;
    leave it out and you get a labelled placeholder to design against.
    """
    if image:
        from PIL import Image
        iw, ih = Image.open(image).size
        scale = min(w / iw, h / ih)
        dw, dh = iw * scale, ih * scale
        s.shapes.add_picture(image, Inches(x + (w - dw) / 2),
                             Inches(y + (h - dh) / 2), Inches(dw), Inches(dh))
        return y + h
    rect(s, x, y, w, h, fill=tint, line=RULE, lw=1.0, radius=0.03)
    line(s, x, y + h, x + w, y + h, color=RULE, lw=1.25)
    line(s, x, y, x, y + h, color=RULE, lw=1.25)
    text(s, x, y + h / 2 - 0.26, w, 0.30, label, size=13, bold=True,
         color=SLATE, align=PP_ALIGN.CENTER, font=HEAD)
    if hint:
        text(s, x + 0.2, y + h / 2 + 0.04, w - 0.4, 0.44, hint, size=10,
             color=SLATE, align=PP_ALIGN.CENTER, italic=True)
    return y + h


def plot_grid(s, x, y, w, h, labels, cols=2, gap=0.22, hints=None,
              images=None):
    """Small multiples: a grid of holders. Pass images=[...] to fill them."""
    n = len(labels)
    rows = (n + cols - 1) // cols
    cw = (w - gap * (cols - 1)) / cols
    ch = (h - gap * (rows - 1)) / rows
    for i, lab in enumerate(labels):
        r, c = divmod(i, cols)
        plot_holder(s, x + c * (cw + gap), y + r * (ch + gap), cw, ch, lab,
                    hints[i] if hints else None,
                    images[i] if images else None)
    return y + h


def chart_with_readout(s, y, chart_label, findings, x=M, w=None, h=4.3,
                       split=0.62, hint=None, image=None, head="What it says",
                       accent=PINE):
    """
    The default analysis layout: the evidence on the left, the claim it
    supports on the right. Never show a chart without its readout.
    """
    w = w or CW
    cwid = w * split - 0.14
    plot_holder(s, x, y, cwid, h, chart_label, hint, image)
    note(s, x + cwid + 0.28, y, w - cwid - 0.28, h, head, findings,
         accent=accent)
    return y + h


# ---------------------------------------------------------------------------
# COMPONENT -- tables
# ---------------------------------------------------------------------------
def table(s, x, y, w, rows, col_w=None, head_fill=NAVY, size=11.5,
          head_size=11.5, row_h=0.34, head_h=0.38, align=None, zebra=True,
          highlight=None, highlight_col=None, highlight_tint=None):
    """
    rows[0] is the header. col_w = relative widths. align = list of
    "l"/"c"/"r". highlight = row index (1-based into the body) to tint;
    highlight_col does the same for a column.
    """
    ncol = len(rows[0])
    nrow = len(rows)
    shape = s.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(w),
                               Inches(head_h + (nrow - 1) * row_h))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False

    col_w = col_w or [1] * ncol
    tot = sum(col_w)
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Emu(int(Inches(w) * cw / tot))
    tbl.rows[0].height = Inches(head_h)
    for r in range(1, nrow):
        tbl.rows[r].height = Inches(row_h)

    amap = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
    align = align or (["l"] + ["r"] * (ncol - 1))

    for r in range(nrow):
        for c in range(ncol):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.10)
            cell.margin_right = Inches(0.10)
            cell.margin_top = Inches(0.035)
            cell.margin_bottom = Inches(0.035)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = head_fill
            elif highlight_col is not None and c == highlight_col:
                cell.fill.fore_color.rgb = TINT_PINE
            elif highlight is not None and r == highlight:
                cell.fill.fore_color.rgb = highlight_tint or TINT_CLAY
            elif zebra and r % 2 == 0:
                cell.fill.fore_color.rgb = TINT_GREY
            else:
                cell.fill.fore_color.rgb = WHITE

            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = amap[align[c] if r else ("l" if c == 0 else
                                                   align[c])]
            val = rows[r][c]
            run = p.add_run()
            run.text = str(val)
            f = run.font
            f.size = Pt(head_size if r == 0 else size)
            f.name = BODY
            f.bold = (r == 0) or (c == 0 and ncol > 2)
            f.color.rgb = WHITE if r == 0 else BLACK
    return y + head_h + (nrow - 1) * row_h


# ---------------------------------------------------------------------------
# COMPONENT -- flow / process
# ---------------------------------------------------------------------------
def flow_box(s, x, y, w, h, head, sub=None, fill=TINT_NAVY, accent=NAVY,
             size=11.5, sub_size=9.5):
    rect(s, x, y, w, h, fill=fill, line=accent, lw=1.25, radius=0.09)
    if sub:
        text(s, x + 0.09, y + 0.11, w - 0.18, h * 0.44, head, size=size,
             bold=True, color=accent, align=PP_ALIGN.CENTER, font=HEAD,
             line_spacing=0.94)
        text(s, x + 0.09, y + h * 0.52, w - 0.18, h * 0.44, sub,
             size=sub_size, color=SLATE, align=PP_ALIGN.CENTER,
             line_spacing=0.94)
    else:
        text(s, x + 0.09, y, w - 0.18, h, head, size=size, bold=True,
             color=accent, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             font=HEAD, line_spacing=0.94)
    return x + w


def flow_row(s, x, y, w, h, steps, gap=0.42, accent=NAVY, fill=TINT_NAVY,
             arrows=True, sub_size=9.5):
    """steps = [(head, sub), ...] laid left to right with arrows between."""
    n = len(steps)
    bw = (w - gap * (n - 1)) / n
    for i, st in enumerate(steps):
        head, sub = (st if isinstance(st, (tuple, list)) else (st, None))
        bx = x + i * (bw + gap)
        acc = accent[i] if isinstance(accent, list) else accent
        fl = fill[i] if isinstance(fill, list) else fill
        flow_box(s, bx, y, bw, h, head, sub, fill=fl, accent=acc,
                 sub_size=sub_size)
        if arrows and i < n - 1:
            arrow(s, bx + bw + 0.07, y + h / 2, bx + bw + gap - 0.07,
                  y + h / 2, color=SLATE)
    return y + h


def numbered_steps(s, x, y, w, steps, size=12.5, gap=0.14, dia=0.38,
                   accent=NAVY):
    """steps = [(title, detail), ...] as a vertical numbered list."""
    cy = y
    for i, (head, detail) in enumerate(steps, 1):
        rect(s, x, cy, dia, dia, fill=accent, radius=0.5)
        text(s, x, cy, dia, dia, str(i), size=12, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=HEAD)
        tw = w - dia - 0.18
        text(s, x + dia + 0.18, cy - 0.015, tw, 0.26, head, size=size,
             bold=True, color=accent, font=HEAD)
        lines = _wrapped_lines(detail, tw, size - 1.5)
        text(s, x + dia + 0.18, cy + 0.255, tw, lines * 0.185 + 0.06, detail,
             size=size - 1.5, color=SLATE, line_spacing=0.98)
        cy += max(dia, 0.255 + lines * 0.185) + gap
    return cy


# ---------------------------------------------------------------------------
# COMPONENT -- claim / test / verdict
# ---------------------------------------------------------------------------
def verdict_block(s, x, y, w, claim, test, verdict, passed=None, h=1.62,
                  gap=0.24):
    """
    Claim -> Test -> Verdict. The unit of an honest research slide, and the
    thing an interviewer is actually listening for.
    """
    cw = (w - 2 * gap) / 3
    acc = CLAY if passed is False else (PINE if passed else NAVY)
    card(s, x, y, cw, h, "Claim", claim, tint=TINT_NAVY, accent=NAVY,
         bullet_body=False, size=11.5)
    card(s, x + cw + gap, y, cw, h, "How it was tested", test,
         tint=TINT_GREY, accent=SLATE, bullet_body=False, size=11.5)
    card(s, x + 2 * (cw + gap), y, cw, h, "Verdict", verdict,
         tint={CLAY: TINT_CLAY, PINE: TINT_PINE, NAVY: TINT_NAVY}[acc],
         accent=acc, bullet_body=False, size=11.5)
    for i in range(2):
        arrow(s, x + (i + 1) * cw + i * gap + 0.045, y + h / 2,
              x + (i + 1) * (cw + gap) - 0.045, y + h / 2, color=SLATE)
    return y + h


def compare(s, x, y, w, h, left, right, left_head="Before",
            right_head="After", gap=0.30, left_accent=SLATE,
            right_accent=PINE):
    """Two-column before/after. Body items are lists of strings."""
    cw = (w - gap) / 2
    card(s, x, y, cw, h, left_head, left, tint=TINT_GREY, accent=left_accent)
    card(s, x + cw + gap, y, cw, h, right_head, right,
         tint={PINE: TINT_PINE, CLAY: TINT_CLAY,
               NAVY: TINT_NAVY}.get(right_accent, TINT_PINE),
         accent=right_accent)
    return y + h


# ---------------------------------------------------------------------------
# Title and section slides
# ---------------------------------------------------------------------------
def title_slide(d, head, sub, meta=None, motif=True):
    s = d.slide(bg=INK)
    text(s, M, 2.30, CW * 0.86, 1.10, head, size=42, bold=True, color=WHITE,
         font=HEAD)
    text(s, M, 3.48, CW * 0.78, 0.80, sub, size=16, color=RGBColor(
        0xB9, 0xC4, 0xE0), line_spacing=1.12)
    if motif:
        timebar(s, M, 4.60, CW * 0.52, 0.14,
                [(3, NAVY, ""), (1, PINE, ""), (1, CLAY, "")], radius=0.4)
    if meta:
        text(s, M, 5.10, CW * 0.8, 0.90, meta, size=12,
             color=RGBColor(0x8C, 0x98, 0xB8), line_spacing=1.20)
    return s


def section_slide(d, number, head, blurb=None, total=6, index=None):
    """
    Divider. The progress motif is the same timebar used for the split
    geometry -- the deck's one repeated visual idea.
    """
    s = d.slide(bg=INK)
    text(s, M, 2.72, 1.30, 0.70, f"{number:02d}", size=44, bold=True,
         color=PINE, font=HEAD)
    text(s, M + 1.32, 2.80, CW - 1.4, 0.66, head, size=33, bold=True,
         color=WHITE, font=HEAD)
    if blurb:
        text(s, M + 1.32, 3.52, CW * 0.66, 0.60, blurb, size=13,
             color=RGBColor(0xA9, 0xB6, 0xD4), line_spacing=1.10)
    idx = index if index is not None else number
    segs = []
    for i in range(1, total + 1):
        segs.append((1, PINE if i == idx else RGBColor(0x2A, 0x33, 0x5C), ""))
        if i < total:
            segs.append((0.10, None, ""))
    timebar(s, M, 4.62, CW * 0.60, 0.11, segs, radius=0.4)
    return s


def closing_slide(d, head, sub=None, meta=None):
    s = d.slide(bg=INK)
    text(s, M, 2.72, CW, 1.00, head, size=44, bold=True, color=WHITE,
         font=HEAD, align=PP_ALIGN.CENTER)
    if sub:
        text(s, M, 3.80, CW, 0.44, sub, size=17,
             color=RGBColor(0xB9, 0xC4, 0xE0), align=PP_ALIGN.CENTER)
    if meta:
        text(s, M, 4.42, CW, 0.40, meta, size=12,
             color=RGBColor(0x8C, 0x98, 0xB8), align=PP_ALIGN.CENTER,
             font=MONO)
    return s
